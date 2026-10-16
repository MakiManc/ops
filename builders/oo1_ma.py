#!/usr/bin/env python3
"""OO1 from Finance's management-account summary (Ross, 15/10/2026).

THE OBJECTIVE, VERBATIM: "OO1: 7% Net Profit", with
  KR1: Actual W/R% (20%)    wages over sales
  KR2: Actual FR% (35%)     food costs over sales
  KR3: Actual V/R% (13%)    variable costs over sales
  KR4: Actual NP% (4-10%)   net profit over sales
  KR5: MA Completion by 7th of every month

ROSS'S DECISIONS, 15/10/2026: "use this sheet to fill OO1 pull on the 3rd of
every month", then, answering the questions that raised:
  * the sheet is Finance's 'AA MA Summary.xlsx' (its 2026 tab), not the
    13/10 copy of it;
  * each KR is PASS/FAIL ON ITS OWN TARGET - 100 met, 0 missed (OKR_BANDS
    oo1_wr / oo1_fr / oo1_vr / oo1_np / oo1_ma). Net profit passes at 4% or more;
  * KR5: "Just pull on the 7th" - the workbook is read on the 7th of every
    month, and KR5 is whether the month just gone was complete in it;
  * it REPLACES the Operations Input sheet's '2026 Summary' block as OO1's
    source.

THE PUBLICATION RULE, WHICH IS THE POINT OF THIS MODULE'S SHAPE.
MakiManc/ops is public and the workbook is Finance's: every £ and every % in it
is a Finance figure. They are read into memory on the runner, compared with the
KR's target, and dropped. What this module writes is, per month and KR,
whether the target was met and the score, and per month which of Finance's
lines were not entered yet - never a figure, a ratio, or a margin from the
target. doc_problems() checks every document against that shape before it is
written, and the bake and the verifier run the same check on what they read.
The workbook's Drive ID is a repository secret of the private ETL repo
(MA_FILE_ID), never a literal and never in this repo: the file is link-shared,
so its ID works as a password. Nothing here prints it, and no exception text
(which can carry the request URL) is ever printed.

HOW A MONTH IS READ. The 2026 tab is a block of rows labelled in column A
(Restaurant Sales, Money In, Wages, Rent, Food costs, Variable costs, TAX,
Money Out, Net Profit, and Finance's own percentage rows) against a header row
of month names Jan..Dec. Labels are matched trimmed and case-blind, and the
month columns are found by their header, so a column added to the right
moves nothing.
  * ENTERED means a number typed into one of Finance's six input lines (sales,
    wages, rent, food, variable, tax). A blank is NOT ENTERED, never 0. The
    formula rows (Money In, Money Out, Net Profit, the percentage rows) read 0
    or #DIV/0! for a month not entered, so they never decide that.
  * UNREADABLE means an input cell that is not a typed number and not a typed
    blank: text, a dash, an error, or a FORMULA showing 0 or nothing (which
    may be a real 0 or a month nobody entered - the cell cannot say). It is
    named, never scored, and never proves KR5 either way.
  * Each KR needs its own lines: KR1 sales and wages, KR2 sales and food, KR3
    sales and variable costs, KR4 all six (net profit takes every cost line,
    tax included). A month missing any of them is not scored for that KR, and
    says which line.
  * Each ratio is computed from the money lines and CHECKED against Finance's
    own rows (its percentage row; for KR4 also its Money Out and Net Profit):
    if they disagree the month is not scored, because one of the two is wrong
    and this module cannot tell which.
  * Only a month that has ENDED before the pull is read.
  * COMPLETE (KR5) means all six input lines entered.

KR5, AND WHY THE PULL TIME MATTERS. KR5's deadline for a month is the 7th of
the next month (UK date). Every verdict keeps the pull that decided it:
  * a pull on or before the deadline that finds the month complete -> 100;
  * a pull AFTER the deadline that finds a line blank -> 0 (an entry is not
    un-typed) - unless an earlier pull saw the month complete, when a line was
    cleared since and the blank proves nothing;
  * a pull ON the deadline day that finds a line blank -> provisional: not
    scored, because Finance may still finish that day. Only the FIRST pull
    after the deadline day decides it: a line blank in both is a final 0
    (naming only those lines); none blank in both proves nothing, and the
    verdict is dropped. The calendar never finalises it (review, 15/10/2026);
  * a pull after the deadline that finds it complete proves nothing about the
    7th, so the month stays unscored for KR5;
  * an UNREADABLE line proves nothing either way.

WHEN IT PULLS (due()). The bake runs after each daily export, at no fixed
hour, and some days not at all. So:
  * a bake on the 7th (UK) always pulls - a second bake that day re-reads;
  * after the 7th, a bake pulls while no pull on or after this month's 7th has
    succeeded (the catch-up: a day with no bake or a failed read loses nothing);
  * on ANY day, a bake pulls while a deadline-day blank waits for the pull
    that decides it.
Any other day it does nothing. `--force` (the bake's dispatch input) pulls
whatever the day.

Run (the ETL bake runs this): MA_FILE_ID=... GOOGLE_SA_JSON=... \
    python3 builders/oo1_ma.py [--force]
Test without Google: python3 builders/oo1_ma.py --from-xlsx F --now ISO --force
"""
from __future__ import annotations

import argparse
import datetime
import io
import json
import math
import os
import re
import sys

#: The tab this reads: the OKRs are 2026's.
MA_YEAR = "2026"
#: Where the workbook lives, as the basis strings name it. The file NAME only:
#: its ID is in the private ETL repo.
MA_SOURCE = "Finance's management-account summary ('AA MA Summary.xlsx', 2026 tab)"
#: KR5's deadline, and the pull day: the 7th of the month after.
MA_DEADLINE_DAY = 7
#: Who decided, for the basis strings.
MA_DECIDED = "Ross, 15/10/2026"
#: The files this writes, next to the snapshot (data/ops_command/).
MA_FILE = "oo1_ma.json"
MA_STATUS_FILE = "oo1_ma_status.json"
#: A workbook this large is not the summary.
MA_MAX_BYTES = 25 * 1024 * 1024
DRIVE_FILES = "https://www.googleapis.com/drive/v3/files"
SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]

_MONTH_ABBR = ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep",
               "oct", "nov", "dec")

#: Column-A label (trimmed, lower-case, inner spaces collapsed) -> key.
ROW_LABELS = {
    "restaurant sales": "sales",
    "money in (£)": "money_in",
    "wages": "wages",
    "wages %": "wages_pct",
    "rent": "rent",
    "food costs": "food",
    "food costs %": "food_pct",
    "variable costs": "variable",
    "variable costs %": "variable_pct",
    "tax": "tax",
    "money out (£)": "money_out",
    "net profit (£)": "np",
    "net profit (%)": "np_pct",
}
#: Finance's six input lines, as the basis names them. A month is COMPLETE
#: when every one is entered.
INPUTS = (("sales", "Restaurant Sales"), ("wages", "Wages"), ("rent", "Rent"),
          ("food", "Food costs"), ("variable", "Variable costs"), ("tax", "TAX"))
INPUT_NAME = dict(INPUTS)
#: The rows that must be on the tab for anything to be read.
REQUIRED_ROWS = ("sales", "money_in", "wages", "rent", "food", "variable", "tax",
                 "money_out", "np", "wages_pct", "food_pct", "variable_pct",
                 "np_pct")

#: kr -> (cost line or None for net profit, Finance's percentage row, band,
#:        what the basis calls it, the input lines needed)
KRS = {
    "KR1": ("wages", "wages_pct", "oo1_wr", "wages", ("sales", "wages")),
    "KR2": ("food", "food_pct", "oo1_fr", "food costs", ("sales", "food")),
    "KR3": ("variable", "variable_pct", "oo1_vr", "variable costs",
            ("sales", "variable")),
    "KR4": (None, "np_pct", "oo1_np", "net profit", tuple(k for k, _ in INPUTS)),
}
#: KR5's band: 1 = complete by the deadline, 0 = not.
KR5_BAND = "oo1_ma"

#: How close a recomputed ratio must be to Finance's own percentage cell (as
#: a fraction: 0.0001 is 0.01 percentage points).
PCT_TOL = 0.0001
#: ...and a recomputed money total to Finance's own (£1 of rounding).
GBP_TOL = 1.0

#: Why a pull read nothing, in words a reader can act on. Never a figure.
CAUSE_TEXT = {
    "unreadable": "the file could not be opened as an .xlsx workbook",
    "tab_missing": "the workbook has no 2026 tab",
    "tab_ambiguous": "more than one tab could be the 2026 tab",
    "header_missing": "the 2026 tab has no Jan..Dec header row",
    "rows_missing": "the 2026 tab is missing a row this reads (a line renamed or removed)",
    "rows_duplicated": "a row this reads appears more than once on the 2026 tab",
    "no_file_id": "MA_FILE_ID is not set in the bake workflow",
    "bad_file_id": "MA_FILE_ID is not a Drive file ID - it should be only the characters "
                   "between /d/ and the next / in the workbook's Drive link",
    "no_credentials": "the bake has no Google service-account key (GOOGLE_SA_JSON)",
    "not_shared": "Drive refused the service account (HTTP 403) - share the workbook "
                  "with the service account as Viewer",
    "not_xlsx": "the file at MA_FILE_ID is a Google Sheet, not the .xlsx workbook (HTTP 403) "
                "- MA_FILE_ID should be the .xlsx file's ID",
    "download_blocked": "the workbook's owner has turned off downloading for viewers (HTTP "
                        "403), so the service account can open it but not read it",
    "rate_limited": "Drive asked the bake to slow down (HTTP 403/429) - the next bake retries",
    "drive_api_disabled": "the Drive API is not enabled on the service account's Google "
                          "Cloud project (HTTP 403) - enable it in the Cloud console",
    "not_found": "Drive has no file at that ID that the service account can read (HTTP 404) "
                 "- its sharing was tightened, or it was moved or deleted; share it with "
                 "the service account as Viewer, and check MA_FILE_ID",
    "http_error": "Drive returned an error",
    "network": "the download failed (network or timeout)",
    "too_large": "the download is far larger than the summary workbook",
    "refused": "the result failed the publication check, so nothing was written",
    "no_reader": "the bake could not load openpyxl, the .xlsx reader - the runner's "
                 "fault, not the workbook's",
    "store_unreadable": "the stored oo1_ma.json is unreadable or fails the publication "
                        "check, so nothing was read - restore the last good copy from "
                        "git history (it holds KR5 verdicts no later read can prove)",
}

#: The not-measured reasons judge_month() can write - FIXED TEXT, so
#: doc_problems() can refuse anything else (a free-text field is where a
#: figure would slip through).
_R_NOT_ENTERED = "not_entered"
_R_NO_SALES = "Money In is not a positive number for this month, so no share of sales can be read"
_R_NP_UNREADABLE = ("Finance's Money Out or Net Profit cell for this month is not a number, "
                    "so the net profit cannot be checked")
_R_NP_DISAGREES = ("Finance's Money Out or Net Profit for this month does not equal its own "
                   "cost lines added up, so it is not scored until the workbook agrees with itself")
_R_PCT_DISAGREES = ("Finance's own percentage row for this month does not match its money "
                    "lines, so it is not scored until the workbook agrees with itself")
FIXED_REASONS = (_R_NO_SALES, _R_NP_UNREADABLE, _R_NP_DISAGREES, _R_PCT_DISAGREES)


#: An input cell that is neither a typed number nor a typed blank: text, a
#: dash, an error, or a formula showing 0 or nothing. NOT the same as blank:
#: it is not "not entered", so it never proves a KR5 0 or 100 (review,
#: 15/10/2026), and the KRs that need it say so.
UNREADABLE = "unreadable"
#: KR verdict key for a line that is entered but not a number.
_R_UNREADABLE = "unreadable"


def _norm(label) -> str:
    return re.sub(r"\s+", " ", str(label or "").strip()).lower()


def _blank(cell) -> bool:
    return cell is None or (isinstance(cell, str) and not cell.strip())


def _formula_cells(formulas: list) -> set:
    """{(row, col)} 0-based cells that hold a formula, from a data_only=False read.

    A plain formula is a string starting '='. An array or dynamic-array formula
    is an object with .ref/.text, and its WHOLE range is formula output - the
    spilled cells carry no formula of their own (review, 15/10/2026).
    """
    out = set()
    try:
        from openpyxl.utils.cell import range_boundaries
    except Exception:  # noqa: BLE001 - grids built in tests carry no objects
        range_boundaries = None
    for i, r in enumerate(formulas or []):
        for j, f in enumerate(r):
            if isinstance(f, str) and f.startswith("="):
                out.add((i, j))
            elif hasattr(f, "ref") and hasattr(f, "text"):
                out.add((i, j))
                ref = getattr(f, "ref", None)
                if range_boundaries and isinstance(ref, str):
                    try:
                        c1, r1, c2, r2 = range_boundaries(ref)
                        out.update((rr - 1, cc - 1) for rr in range(r1, r2 + 1)
                                   for cc in range(c1, c2 + 1))
                    except Exception:  # noqa: BLE001
                        pass
    return out


def grid_from_xlsx(data: bytes, year: str = MA_YEAR) -> dict:
    """{"title", "rows", "formula_cells"} for the year's tab only, or {"cause"}.

    Only that tab is read - the workbook's other tabs are never loaded into
    rows. It is read TWICE: data_only=True for what Excel or Google last
    calculated (what Finance sees), and data_only=False to know which cells
    are formulas. An input line that is a formula showing 0 or nothing (or
    never calculated) may be a month nobody has entered or a real 0, so it
    reads UNREADABLE - never a typed 0 and never a typed blank (review,
    15/10/2026).
    """
    from openpyxl import load_workbook  # only the pull needs it
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        names = list(wb.sheetnames)
        by_title = [t for t in names if _norm(t) == year]
        if len(by_title) > 1:
            return {"cause": "tab_ambiguous"}
        if not by_title:
            by_a1 = []
            for t in names:
                first = next(wb[t].iter_rows(min_row=1, max_row=1, values_only=True), ())
                if first and _norm(first[0]) == year:
                    by_a1.append(t)
            if len(by_a1) != 1:
                return {"cause": "tab_ambiguous" if by_a1 else "tab_missing"}
            by_title = by_a1
        title = by_title[0]
        ws = wb[title]
        ws.reset_dimensions()  # a file that under-reports its size is read in full
        rows = [list(r) for r in ws.iter_rows(min_row=1, min_col=1, values_only=True)]
    finally:
        wb.close()
    wf = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    try:
        ws = wf[title]
        ws.reset_dimensions()
        formulas = [list(r) for r in ws.iter_rows(min_row=1, min_col=1, values_only=True)]
    finally:
        wf.close()
    # the title is free text and never leaves this function: the pull records
    # the year it matched (review, 15/10/2026)
    return {"title": title, "rows": rows, "formula_cells": _formula_cells(formulas)}


def find_tab(sheets: dict, year: str = MA_YEAR):
    """(title, rows) of the year's tab, or (None, cause).

    The tab whose title is the year (trimmed), else the one tab whose A1 is.
    """
    by_title = [t for t in sheets if _norm(t) == year]
    if len(by_title) == 1:
        return by_title[0], sheets[by_title[0]]
    if len(by_title) > 1:
        return None, "tab_ambiguous"
    by_a1 = [t for t, rows in sheets.items()
             if rows and rows[0] and _norm(rows[0][0]) == year]
    if len(by_a1) == 1:
        return by_a1[0], sheets[by_a1[0]]
    return None, ("tab_ambiguous" if by_a1 else "tab_missing")


def num(cell):
    """A number Finance typed, or None for blank / text / an error cell.

    An error cell ('#DIV/0!') and a blank both read None here: the caller
    decides what that means. A bool is not a number.
    """
    if isinstance(cell, bool) or cell is None:
        return None
    if isinstance(cell, (int, float)):
        return float(cell) if math.isfinite(cell) else None
    s = str(cell).strip().replace(",", "").replace("£", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if re.fullmatch(r"-?\d+(\.\d+)?", s):
        v = float(s)
        return -v if neg else v
    return None


def _mname(c):
    # a header typed as a date and formatted 'mmm' reads as a datetime
    if isinstance(c, (datetime.date, datetime.datetime)):
        return _MONTH_ABBR[c.month - 1]
    return _norm(c)[:3]


def _month_cols(r):
    names = [_mname(c) for c in r]
    for j in range(len(names) - 11):
        if tuple(names[j:j + 12]) == _MONTH_ABBR:
            return {m + 1: j + m for m in range(12)}
    return None


def read_months(rows: list, formula_cells: set | None = None):
    """({'YYYY-MM': {key: value}}, cause or None) for the year's tab.

    A value is a number; None for a BLANK input (not entered) or an unreadable
    formula row; UNREADABLE for an input cell holding text, a dash or an error,
    or - by formula_cells, from grid_from_xlsx - a formula showing 0 or nothing
    (a real 0 and a month nobody entered look the same there).

    The label block is found first; its header is the LAST row holding Jan..Dec
    in twelve consecutive cells ABOVE the first label (a full month name, or a
    date formatted 'mmm', counts) - so a strip with its own month header higher
    up the tab cannot shift every month (review, 15/10/2026). Each month's
    column is where its name sits. A required row missing, or one appearing
    twice, is a changed layout: nothing is read rather than the wrong row.
    """
    first_label = next((i for i, r in enumerate(rows)
                        if r and ROW_LABELS.get(_norm(r[0])) is not None), None)
    if first_label is None:
        return None, ("header_missing" if not any(_month_cols(r) for r in rows) else "rows_missing")
    header = cols = None
    for i in range(first_label - 1, -1, -1):
        c = _month_cols(rows[i])
        if c:
            header, cols = i, c
            break
    if header is None:
        return None, "header_missing"
    found, at = {}, {}
    for i, r in enumerate(rows[header + 1:], start=header + 1):
        key = ROW_LABELS.get(_norm(r[0] if r else ""))
        if key is None:
            continue
        if key in found:
            return None, "rows_duplicated"
        found[key], at[key] = r, i
    if any(k not in found for k in REQUIRED_ROWS):
        return None, "rows_missing"
    inputs = set(INPUT_NAME)
    fc = formula_cells or set()

    def _cell(k, r, j):
        raw = r[j] if j < len(r) else None
        if k not in inputs:
            return num(raw)
        if (at[k], j) in fc and (_blank(raw) or num(raw) == 0):
            return UNREADABLE   # a formula's 0 or '': not entered, or a real 0
        if _blank(raw):
            return None
        v = num(raw)
        return v if v is not None else UNREADABLE
    out = {}
    for m, j in cols.items():
        out[f"{MA_YEAR}-{m:02d}"] = {k: _cell(k, r, j) for k, r in found.items()}
    return out, None


def deadline(ym: str) -> datetime.date:
    """KR5's deadline for month ym: the 7th of the month after."""
    y, m = int(ym[:4]), int(ym[5:7])
    return datetime.date(y + (m == 12), 1 if m == 12 else m + 1, MA_DEADLINE_DAY)


def judge_month(v: dict, okr_score) -> dict:
    """One ended month's KR1-4 verdicts and completeness. No figures out.

    okr_score is the bake's own (OKR_BANDS is the only place a threshold
    lives). Returns {"KR1": {"score", "met", "checked"} |
    {"not_measured": fixed text} | {"not_entered": [names]} |
    {"unreadable": [names]}, ..., "complete": bool, "missing": [names],
    "unreadable": [names]}.
    """
    missing = [name for k, name in INPUTS if v.get(k) is None]
    unread = [name for k, name in INPUTS if v.get(k) == UNREADABLE]
    res = {"complete": not missing and not unread, "missing": missing, "unreadable": unread}
    mi = v.get("money_in")
    for kr, (line, pct_row, band, what, needs) in KRS.items():
        gone = [INPUT_NAME[k] for k in needs if v.get(k) is None]
        if gone:
            res[kr] = {_R_NOT_ENTERED: gone}
            continue
        bad = [INPUT_NAME[k] for k in needs if v.get(k) == UNREADABLE]
        if bad:
            res[kr] = {_R_UNREADABLE: bad}
            continue
        if mi is None or mi <= 0:
            res[kr] = {"not_measured": _R_NO_SALES}
            continue
        if line is None:
            out_ = sum(v[k] for k, _ in INPUTS if k != "sales")
            fin_out, fin_np = v.get("money_out"), v.get("np")
            if fin_out is None or fin_np is None:
                res[kr] = {"not_measured": _R_NP_UNREADABLE}
                continue
            if abs(fin_out - out_) > GBP_TOL or abs(fin_np - (mi - out_)) > GBP_TOL:
                res[kr] = {"not_measured": _R_NP_DISAGREES}
                continue
            ratio = (mi - out_) / mi
        else:
            ratio = v[line] / mi
        fin_pct = v.get(pct_row)
        if fin_pct is not None and abs(fin_pct - ratio) > PCT_TOL:
            res[kr] = {"not_measured": _R_PCT_DISAGREES}
            continue
        # The band works in percent, as the KR is written. Rounded to 6dp so
        # binary noise cannot push an exact 20% over the line.
        s = okr_score(band, round(100.0 * ratio, 6))
        res[kr] = {"score": s, "met": s == 100, "checked": fin_pct is not None}
    return res


def judge_pull(sheets_or_bytes, pulled_at: datetime.datetime, okr_score) -> dict:
    """One pull's public result: {ok, cause, pulled_at, pull_date, tab, months}.

    pulled_at is the pull's UK wall-clock time (the deadline is a UK date).
    Only months that ENDED before pulled_at's date are judged. `tab` is the
    year matched (MA_YEAR), never the tab's own title, which is free text.
    """
    base = {"pulled_at": pulled_at.isoformat(timespec="minutes"),
            "pull_date": pulled_at.date().isoformat()}
    fc = None
    if isinstance(sheets_or_bytes, (bytes, bytearray)):
        try:
            g = grid_from_xlsx(bytes(sheets_or_bytes))
        except ImportError:
            return {**base, "ok": False, "cause": "no_reader", "tab": None}
        except Exception:  # noqa: BLE001 - any unreadable file is one cause
            return {**base, "ok": False, "cause": "unreadable", "tab": None}
        if g.get("cause"):
            return {**base, "ok": False, "cause": g["cause"], "tab": None}
        rows, fc = g["rows"], g["formula_cells"]
    else:  # a {title: rows} grid (tests)
        title, rows = find_tab(sheets_or_bytes)
        if title is None:
            return {**base, "ok": False, "cause": rows, "tab": None}
    months, cause = read_months(rows, fc)
    if cause:
        return {**base, "ok": False, "cause": cause, "tab": MA_YEAR}
    cur = pulled_at.date().strftime("%Y-%m")
    judged = {ym: judge_month(v, okr_score) for ym, v in months.items() if ym < cur}
    return {**base, "ok": True, "cause": None, "tab": MA_YEAR, "months": judged}


def kr5_from_pull(ym: str, res: dict, pulled_at: datetime.datetime, seen_complete=False,
                  seen_entered=None):
    """KR5's verdict for month ym from ONE pull, or None when it proves nothing.

    seen_entered: Finance's line names any EARLIER read saw not blank (a number,
    or unreadable). {"score", "final", "judged_at", "missing"}:
      complete, read on or before the deadline        -> 100, final
      a line BLANK, read on the deadline day          -> 0, provisional: only
          the first read after the deadline decides it (see merge)
      a line BLANK, read after the deadline, that NO earlier read saw entered
          -> 0, final, naming only those lines (an entry is not un-typed - but
          a line seen entered once and blank now was cleared, and proves
          nothing; review, 15/10/2026)
      anything else (before the deadline and incomplete; complete but read
      after it; only unreadable lines)                -> None
    """
    d, day = deadline(ym), pulled_at.date()
    at = pulled_at.isoformat(timespec="minutes")
    complete, missing = bool(res.get("complete")), list(res.get("missing") or [])
    if complete:
        return {"score": 100, "final": True, "judged_at": at, "missing": []} if day <= d else None
    if not missing or day < d:
        return None
    if day == d:
        return {"score": 0, "final": False, "judged_at": at, "missing": missing}
    prior = set(seen_entered if seen_entered is not None
                else (_INPUT_ORDER if seen_complete else []))
    blank = [n for n in missing if n not in prior]
    return {"score": 0, "final": True, "judged_at": at, "missing": blank} if blank else None


#: Why a month's KR5 was decided as NOT JUDGED (the bake words them):
#:   after - no read judged the month on or before its deadline, and the first
#:           read after it found no line blank;
#:   late  - a read on or before the deadline found it incomplete, and the
#:           first read after it found no line that every read had seen blank.
KR5_NOT_JUDGED = ("after", "late")


def merge(doc: dict | None, pull: dict) -> dict:
    """Fold one pull into the stored document (the public file).

    KR1-4 for a month come from the NEWEST successful pull that judged it
    (Finance can correct a month; the latest workbook is the record). The
    month also remembers which of Finance's lines any read saw not blank
    (seen_entered) and whether any read saw it complete (seen_complete).
    KR5 (review, 15/10/2026):
      * on or before the deadline: a complete read is a final 100; a blank on
        the deadline day is a provisional 0 (the day's latest read stands; a
        later read that day with no blank line drops it);
      * the FIRST successful read after the deadline decides the month FOR
        GOOD: a final 0 naming the lines no read ever saw entered (keeping the
        deadline-day read as provisional_at / provisional_missing), or a
        recorded "not judged" (KR5_NOT_JUDGED) - so a later read can never turn
        a month into a 0 the deciding read did not prove;
      * a decided month (100, final 0, not judged) is never changed.
    Every pull, failed or not, is listed - nothing is dropped.
    """
    doc = {"version": 1, "months": {}, "kr5": {}, "pulls": [], **(doc or {})}
    doc["months"] = dict(doc.get("months") or {})
    doc["kr5"] = dict(doc.get("kr5") or {})
    doc["pulls"] = list(doc.get("pulls") or []) + [
        {k: pull.get(k) for k in ("pulled_at", "pull_date", "ok", "cause", "tab")}]
    if not pull.get("ok"):
        return doc
    at = datetime.datetime.fromisoformat(pull["pulled_at"])
    for ym, res in (pull.get("months") or {}).items():
        prev = doc["months"].get(ym) or {}
        seen_c = bool(prev.get("seen_complete"))
        seen_e = set(prev.get("seen_entered") or (_INPUT_ORDER if seen_c else []))
        now_e = {n for n in _INPUT_ORDER if n not in (res.get("missing") or [])}
        doc["months"][ym] = {**res, "pulled_at": pull["pulled_at"],
                             "seen_complete": seen_c or bool(res.get("complete")),
                             "seen_entered": [n for n in _INPUT_ORDER if n in seen_e | now_e]}
        old = doc["kr5"].get(ym)
        if old and (old.get("final") or "not_judged" in old):
            continue   # decided, for good
        new = kr5_from_pull(ym, res, at, seen_entered=seen_e)
        d = deadline(ym)
        if at.date() <= d:
            if new is not None:
                doc["kr5"][ym] = new          # a 100, or the deadline day's latest blank
            elif at.date() == d and old:
                del doc["kr5"][ym]            # a later read that day shows no blank line
            continue
        prov = ({"provisional_at": old["judged_at"], "provisional_missing": list(old["missing"])}
                if old else {})
        if new is not None:
            doc["kr5"][ym] = {**new, **prov}
        else:
            doc["kr5"][ym] = {"not_judged": "late" if prev else "after",
                              "judged_at": pull["pulled_at"],
                              **({"last_before": prev["pulled_at"]} if prev else {}), **prov}
    return doc


def kr5_now(v: dict | None, ym: str, today: datetime.date):
    """A stored KR5 verdict as of `today` - unchanged. A provisional 0 is NOT
    made final by the calendar: only a read after the deadline day confirms it
    (review, 15/10/2026), and due() asks for that read. Kept for callers."""
    return v


def kr5_awaiting(v: dict | None, ym: str, today: datetime.date) -> bool:
    """A provisional 0 whose deadline day is over: waiting for the read that
    decides it."""
    return (isinstance(v, dict) and "not_judged" not in v and v.get("final") is False
            and today > deadline(ym))


def last_seventh(today: datetime.date) -> datetime.date:
    """The latest 7th on or before `today` (this month's from the 7th, else
    last month's)."""
    if today.day >= MA_DEADLINE_DAY:
        return today.replace(day=MA_DEADLINE_DAY)
    first = today.replace(day=1) - datetime.timedelta(days=1)
    return first.replace(day=MA_DEADLINE_DAY)


def due(doc: dict | None, now: datetime.datetime) -> bool:
    """Whether a bake at UK time `now` should pull (see WHEN IT PULLS): every
    bake on the 7th; any bake while a deadline-day blank waits for its
    deciding read; otherwise while no read since the latest 7th has
    SUCCEEDED (a missed or failed 7th is caught up by the next bake, even
    after the month turns - review, 15/10/2026).
    """
    if now.day == MA_DEADLINE_DAY:
        return True
    if any(kr5_awaiting(v, ym, now.date())
           for ym, v in ((doc or {}).get("kr5") or {}).items()
           if isinstance(ym, str) and _MONTH_RE.fullmatch(ym)):
        return True
    since = last_seventh(now.date()).isoformat()
    return not any(isinstance(p, dict) and p.get("ok") is True
                   and str(p.get("pull_date") or "") >= since
                   for p in (doc or {}).get("pulls") or [])


# ---------------------------------------------------------------------------
# THE PUBLICATION CHECK
# ---------------------------------------------------------------------------
_MONTH_RE = re.compile(r"20\d\d-(0[1-9]|1[0-2])")
_STAMP_RE = re.compile(r"20\d\d-\d\d-\d\dT\d\d:\d\d([+-]\d\d:\d\d)?")
_DATE_RE = re.compile(r"20\d\d-\d\d-\d\d")
_INPUT_NAMES = frozenset(INPUT_NAME.values())
_INPUT_ORDER = [name for _k, name in INPUTS]
#: The percentages an OO1 text may name: the KRs' own targets.
TARGET_PCTS = frozenset({"20%", "35%", "13%", "4%"})
_PCT_RE = re.compile(r"(?<![\w.,·])(-?\d+(?:[.,·]\d+)?)\s*(%|％|per\s?cent\b|percent\b|pc\b)", re.I)
_MONEY_RE = re.compile(r"£|₤|\bGBP\b|\bEUR\b|€|\$|\bpounds?\b|\d[\d,]*\.\d+|\d{1,3}(?:,\d{3})+", re.I)


def figure_in(text) -> bool:
    """Whether an OO1 text carries a figure: a money sign or word, a decimal, a
    thousands-separated amount, or any percentage that is not exactly one of the
    KRs' own targets (a negative or a decimal never matches a target)."""
    t = str(text or "")
    if _MONEY_RE.search(t):
        return True
    for n, unit in _PCT_RE.findall(t):
        if unit != "%" or f"{n}%" not in TARGET_PCTS:
            return True
    return False


def _stamp(x) -> bool:
    """A real minute-resolution timestamp, exactly as this module writes one."""
    if not (isinstance(x, str) and _STAMP_RE.fullmatch(x)):
        return False
    try:
        return datetime.datetime.fromisoformat(x).isoformat(timespec="minutes") == x
    except ValueError:
        return False


def _names(x, allow_empty=True) -> bool:
    """A list of Finance's line names, each once, in the workbook's order."""
    if not isinstance(x, list) or (not x and not allow_empty):
        return False
    if not all(isinstance(n, str) and n in _INPUT_NAMES for n in x) or len(set(x)) != len(x):
        return False
    return x == [n for n in _INPUT_ORDER if n in x]


def doc_problems(doc) -> list:
    """Why `doc` must not be published, or [] - names and paths, never values.

    The document must be EXACTLY the shape merge() writes: every key known,
    every score 0 or 100, every flag a bool, every timestamp a real one, every
    line-name list Finance's own names in order, every reason one of the fixed
    texts, every month of MA_YEAR, every pull's date its stamp's date. A field
    this does not know, or a number where a score or a date belongs, could be a
    Finance figure, so it is refused rather than dropped. A key is never
    repeated in a message (it is where a figure would hide).
    """
    out = []
    if not isinstance(doc, dict):
        return ["the document is not an object"]
    extra = [k for k in doc if k not in ("version", "months", "kr5", "pulls")]
    if extra:
        out.append(f"{len(extra)} unknown top-level field(s)")
    if type(doc.get("version")) is not int or doc.get("version") != 1:
        out.append("version is not 1")

    def _ym(x):
        return isinstance(x, str) and _MONTH_RE.fullmatch(x) is not None and x[:4] == MA_YEAR

    months = doc.get("months")
    if not isinstance(months, dict):
        out.append("months is not an object")
        months = {}
    for ym, m in months.items():
        if not _ym(ym) or not isinstance(m, dict):
            out.append("months: an entry that is not a month of the year")
            continue
        for k, v in m.items():
            if k in KRS:
                if not isinstance(v, dict) or len(v) == 0:
                    out.append(f"{ym} {k} is not a verdict")
                elif set(v) == {"score", "met", "checked"}:
                    if (type(v["score"]) is not int or v["score"] not in (0, 100)
                            or not isinstance(v["met"], bool) or v["met"] != (v["score"] == 100)
                            or not isinstance(v["checked"], bool)):
                        out.append(f"{ym} {k} is not a 0/100 score with its flags")
                elif set(v) == {"not_measured"}:
                    if v["not_measured"] not in FIXED_REASONS:
                        out.append(f"{ym} {k} carries a reason that is not one of the fixed texts")
                elif set(v) in ({_R_NOT_ENTERED}, {_R_UNREADABLE}):
                    if not _names(next(iter(v.values())), allow_empty=False):
                        out.append(f"{ym} {k} names a line that is not one of Finance's six")
                else:
                    out.append(f"{ym} {k} is not a verdict")
            elif k in ("complete", "seen_complete"):
                if not isinstance(v, bool):
                    out.append(f"{ym} {k} is not true/false")
            elif k in ("missing", "unreadable", "seen_entered"):
                if not _names(v):
                    out.append(f"{ym} {k} names a line that is not one of Finance's six")
            elif k == "pulled_at":
                if not _stamp(v):
                    out.append(f"{ym} pulled_at is not a timestamp")
            else:
                out.append(f"{ym} carries an unknown field")
        if m.get("complete") is True and (m.get("missing") or m.get("unreadable")):
            out.append(f"{ym} is complete with lines missing")
        if isinstance(m.get("seen_entered"), list) and isinstance(m.get("missing"), list) and not (
                (set(_INPUT_ORDER) - set(m["missing"])) <= set(m["seen_entered"])):
            out.append(f"{ym} has lines seen entered that its record does not")
    kr5 = doc.get("kr5")
    if not isinstance(kr5, dict):
        out.append("kr5 is not an object")
        kr5 = {}
    for ym, v in kr5.items():
        ok = _ym(ym) and isinstance(v, dict)
        if ok and "not_judged" in v:
            # a month decided as not judged (merge); its stamps in order
            ok = ({"not_judged", "judged_at"} <= set(v)
                  <= {"not_judged", "judged_at", "last_before", "provisional_at",
                      "provisional_missing"}
                  and v["not_judged"] in KR5_NOT_JUDGED and _stamp(v["judged_at"])
                  and (("last_before" in v) == (v["not_judged"] == "late"))
                  and ("last_before" not in v
                       or (_stamp(v["last_before"]) and v["last_before"] < v["judged_at"]))
                  and (("provisional_at" in v) == ("provisional_missing" in v))
                  and ("provisional_at" not in v
                       or (v.get("last_before") == v["provisional_at"]
                           and _names(v["provisional_missing"], allow_empty=False))))
        elif ok:
            ok = ({"score", "final", "judged_at", "missing"} <= set(v)
                  <= {"score", "final", "judged_at", "missing", "provisional_at",
                      "provisional_missing"}
                  and (("provisional_at" in v) == ("provisional_missing" in v)))
            if ok:
                ok = (type(v["score"]) is int and v["score"] in (0, 100)
                      and isinstance(v["final"], bool)
                      and _stamp(v["judged_at"]) and _names(v["missing"])
                      and not (v["score"] == 0 and v["final"] and not v["missing"])
                      and ("provisional_at" not in v or (
                          _stamp(v["provisional_at"]) and v["score"] == 0 and v["final"]
                          and v["provisional_at"] < v["judged_at"]
                          and _names(v["provisional_missing"], allow_empty=False)
                          and set(v["missing"]) <= set(v["provisional_missing"])))
                      and not (v["score"] == 100 and (v["missing"] or not v["final"])))
        if not ok:
            out.append("kr5: an entry that is not a KR5 verdict")
    pulls = doc.get("pulls")
    if not isinstance(pulls, list):
        out.append("pulls is not a list")
        pulls = []
    for i, p in enumerate(pulls):
        if (not isinstance(p, dict)
                or set(p) != {"pulled_at", "pull_date", "ok", "cause", "tab"}
                or not _stamp(p["pulled_at"])
                or p["pull_date"] != p["pulled_at"][:10]
                or not isinstance(p["ok"], bool)
                or not (p["cause"] is None or p["cause"] in CAUSE_TEXT)
                or (p["ok"] and p["cause"] is not None)
                or not (p["tab"] is None or p["tab"] == MA_YEAR)):
            out.append(f"pulls[{i}] is not a pull record")
    return out


# ---------------------------------------------------------------------------
# THE PULL (run by the private ETL repo's bake, with its credentials)
# ---------------------------------------------------------------------------
def _atomic_write(path: str, text: str) -> None:
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        fh.write(text)
    os.replace(tmp, path)


def _load(path: str):
    """(doc or None, ok). ok False = present but unreadable or refused."""
    try:
        with open(path, encoding="utf-8") as fh:
            doc = json.load(fh)
    except FileNotFoundError:
        return None, True
    except (OSError, ValueError):
        return None, False
    return (doc, True) if not doc_problems(doc) else (None, False)


_ID_RE = re.compile(r"(?:.*?/d/)?([-\w]{25,})(?:[/?#].*)?")


def file_id(raw):
    """MA_FILE_ID as a bare Drive file ID: None when unset, False when it is not
    one. The whole Drive link, or the part after /d/ with '/edit?...' still on
    it, gives the ID (review, 15/10/2026: 'the part after /d/' is easy to take
    too much of, and a 404 would then blame the workbook)."""
    s_ = str(raw or "").strip()
    if not s_:
        return None
    m = _ID_RE.fullmatch(s_)
    return m.group(1) if m else False


def download(file_id: str) -> tuple:
    """(bytes or None, cause or None, http_status or None). Never raises,
    never puts the ID or an exception message in what it returns."""
    sa = os.environ.get("GOOGLE_SA_JSON", "").strip()
    if not sa:
        return None, "no_credentials", None
    try:
        from google.auth.transport.requests import AuthorizedSession
        from google.oauth2 import service_account
        creds = service_account.Credentials.from_service_account_info(
            json.loads(sa), scopes=SCOPES)
        sess = AuthorizedSession(creds)
        # an older link share can need its resource key: the repository
        # secret MA_RESOURCE_KEY beside MA_FILE_ID (empty when unset)
        rk = os.environ.get("MA_RESOURCE_KEY", "").strip()
        r = sess.get(f"{DRIVE_FILES}/{file_id}",
                     params={"alt": "media", "supportsAllDrives": "true"},
                     headers=({"X-Goog-Drive-Resource-Keys": f"{file_id}/{rk}"} if rk else None),
                     timeout=(15, 120), stream=True)
    except Exception:  # noqa: BLE001 - its text can carry the URL, so only the cause
        return None, "network", None
    try:
        if r.status_code in (403, 429):
            # the body names the reason; it is read, never printed
            try:
                body = r.content[:4000].decode("utf-8", "replace")
            except Exception:  # noqa: BLE001
                body = ""
            if re.search(r"accessNotConfigured|SERVICE_DISABLED|has not been used in project", body):
                return None, "drive_api_disabled", r.status_code
            if r.status_code == 429 or re.search(r"[rR]ateLimitExceeded|quotaExceeded", body):
                return None, "rate_limited", r.status_code
            if "fileNotDownloadable" in body:
                return None, "not_xlsx", 403
            if "cannotDownloadFile" in body:
                return None, "download_blocked", 403
            return None, "not_shared", 403
        if r.status_code == 404:
            # Drive answers 404, not 403, when the caller cannot read the file
            return None, "not_found", 404
        if r.status_code != 200:
            return None, "http_error", r.status_code
        buf = io.BytesIO()
        for chunk in r.iter_content(1 << 16):
            buf.write(chunk)
            if buf.tell() > MA_MAX_BYTES:
                return None, "too_large", 200
        return buf.getvalue(), None, 200
    except Exception:  # noqa: BLE001
        return None, "network", None
    finally:
        r.close()


def _okr_score():
    """The bake's okr_score: OKR_BANDS is the one place a threshold lives."""
    here = os.path.dirname(os.path.abspath(__file__))
    if here not in sys.path:
        sys.path.insert(0, here)
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "bake_ops_command", os.path.join(here, "bake_ops_command.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.okr_score


def _write_status(spath, stamp, ok, cause, http=None):
    """oo1_ma_status.json: fixed texts and stamps only. The previous file's
    last_ok_at / failing_since are carried only when they are real stamps."""
    try:
        with open(spath, encoding="utf-8") as fh:
            prev = json.load(fh)
    except (OSError, ValueError):
        prev = None
    prev = prev if isinstance(prev, dict) else {}
    lok = prev.get("last_ok_at") if _stamp(prev.get("last_ok_at")) else None
    fs = prev.get("failing_since") if _stamp(prev.get("failing_since")) else None
    st = {"attempted_at": stamp, "ok": bool(ok), "cause": cause,
          "detail": CAUSE_TEXT.get(cause) if cause else None,
          "http_status": http if isinstance(http, int) and not isinstance(http, bool) else None,
          "last_ok_at": stamp if ok else lok,
          "failing_since": None if ok else (fs or stamp)}
    _atomic_write(spath, json.dumps(st, indent=1, sort_keys=True) + "\n")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--force", action="store_true",
                    help="pull whatever the day (the bake's force_ma_pull input)")
    ap.add_argument("--from-xlsx", help="read this local file instead of Drive (tests)")
    ap.add_argument("--now", help="the UK time to act as (tests), ISO")
    ap.add_argument("--out-dir", default=os.environ.get("OPS_COMMAND_OUT_DIR", "").strip()
                    or os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
                                    "data", "ops_command"))
    a = ap.parse_args(argv)
    from zoneinfo import ZoneInfo
    uk = ZoneInfo("Europe/London")
    now = (datetime.datetime.fromisoformat(a.now) if a.now else datetime.datetime.now(uk))
    now = now.astimezone(uk) if now.tzinfo else now.replace(tzinfo=uk)
    stamp = now.isoformat(timespec="minutes")
    path = os.path.join(a.out_dir, MA_FILE)
    spath = os.path.join(a.out_dir, MA_STATUS_FILE)
    doc, ok = _load(path)
    if not ok:
        # never overwrite a file this cannot read: it may hold KR5 verdicts
        # that no later pull can prove again - but say so in the status, so
        # the rows and the verifier show reads are being skipped (review)
        print(f"[oo1_ma] {MA_FILE} is present but unreadable or fails the publication "
              "check - left untouched, nothing pulled. Restore the last good copy from git "
              "history; do not delete it")
        _write_status(spath, stamp, False, "store_unreadable")
        return 0
    if not (a.force or due(doc, now)):
        nxt = (now.date().replace(day=MA_DEADLINE_DAY) if now.day < MA_DEADLINE_DAY
               else (now.date().replace(day=1) + datetime.timedelta(days=32)).replace(day=MA_DEADLINE_DAY))
        print(f"[oo1_ma] not due: Finance's MA summary is read on the 7th "
              f"(next {nxt:%d/%m/%Y})")
        return 0
    http = None
    try:
        import openpyxl  # noqa: F401 - checked before downloading: a missing reader
        reader_ok = True  # is the runner's fault, never "the file is unreadable"
    except ImportError:
        reader_ok = False
    if not reader_ok:
        data, cause = None, "no_reader"
    elif a.from_xlsx:
        with open(a.from_xlsx, "rb") as fh:
            data, cause = fh.read(), None
    else:
        fid = file_id(os.environ.get("MA_FILE_ID", ""))
        data, cause, http = ((None, "no_file_id", None) if fid is None
                             else (None, "bad_file_id", None) if fid is False
                             else download(fid))
    if data is None:
        pull = {"pulled_at": stamp, "pull_date": now.date().isoformat(), "ok": False,
                "cause": cause, "tab": None}
    else:
        pull = judge_pull(data, now, _okr_score())
    new = merge(doc, pull)
    probs = doc_problems(new)
    if probs:
        # the shape check names fields only; nothing here can carry a figure
        print(f"[oo1_ma] REFUSED: {len(probs)} problem(s) in the result, e.g. {probs[0]} - "
              f"{MA_FILE} left untouched")
        pull, new = {**pull, "ok": False, "cause": "refused"}, None
    else:
        _atomic_write(path, json.dumps(new, indent=1, sort_keys=True) + "\n")
    _write_status(spath, stamp, bool(pull.get("ok")), pull.get("cause"), http)
    if pull.get("ok"):
        ms = sorted((new or {}).get("months") or {})
        print(f"[oo1_ma] read Finance's MA summary ({MA_YEAR} tab): "
              f"{len(ms)} ended month(s) judged" + (f", {ms[0]}..{ms[-1]}" if ms else ""))
    else:
        print(f"[oo1_ma] pull failed: {pull.get('cause')} - "
              f"{CAUSE_TEXT.get(pull.get('cause'), '')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
